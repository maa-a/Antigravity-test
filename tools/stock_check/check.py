"""案件管理表の払出数量を、SAP在庫データの指定倉庫だけで賄えるか確認する。

使い方:
  python3 check.py <案件管理表.xlsx> <シート名> <開始行> <終了行> <SAP在庫.xlsx> [保管場所コード=CX01] [--xlsx 出力.xlsx]

  --xlsx を付けると、品目ごとの在庫数・払出数・払出後残数の一覧Excelも出力する。

判定ルール:
  - 照合は E列の品目コードで行う（F列の品目名は参考）。
  - SAP側で表示できない漢字が「?」になる差は同一商品とみなす（例: 茈 → ?）。
  - 払出後の残数が 0 は問題なし。マイナスのみ「不足」。
"""
import re, sys, unicodedata, warnings
import openpyxl, pandas as pd

warnings.filterwarnings('ignore')
args = sys.argv[1:]
out = None
if '--xlsx' in args:
    i = args.index('--xlsx'); out = args[i + 1]; del args[i:i + 2]
req_path, sheet, start, end, sap_path = args[0], args[1], int(args[2]), int(args[3]), args[4]
loc = args[5] if len(args) > 5 else 'CX01'

norm = lambda s: unicodedata.normalize('NFKC', str(s or '')).replace(' ', '')

def same_name(a, b):
    """SAP側の '?' は任意の1文字として扱う。"""
    a, b = norm(a), norm(b)
    return len(a) == len(b) and all(x == y or y == '?' for x, y in zip(a, b))

ws = openpyxl.load_workbook(req_path, data_only=True)[sheet]
req = {}
for r in range(start, end + 1):
    code, name, qty = ws.cell(r, 5).value, ws.cell(r, 6).value, ws.cell(r, 8).value
    if code in (None, '---') or qty in (None, ''):
        continue
    e = req.setdefault(code, {'name': name, 'qty': 0, 'rows': []})
    e['qty'] += qty; e['rows'].append(r)

sap = pd.read_excel(sap_path, sheet_name='XTab DS_1')
stock = sap[sap['保管場所'] == loc].groupby('品目コード').agg(name=('品目テキスト', 'first'), qty=('在庫数量', 'sum'))

short, zero, missing, name_diff = [], [], [], []
for code, e in req.items():
    if code not in stock.index:
        missing.append((e['rows'], code, e['name'], e['qty'])); continue
    s = stock.loc[code]; rest = int(s.qty) - e['qty']
    if not same_name(e['name'], s['name']): name_diff.append((e['rows'], code, e['name'], s['name']))
    if rest < 0: short.append((e['rows'], code, e['name'], e['qty'], int(s.qty), -rest))
    elif rest == 0: zero.append((e['rows'], code, e['name'], e['qty']))

print(f'対象 {len(req)} 品目 / 保管場所 {loc}')
print(f'不足（マイナス）: {len(short)} 件')
for x in short: print('  行', x[0], x[1], x[2], f'払出 {x[3]} / 在庫 {x[4]} / 不足 {x[5]}')
print(f'{loc} に在庫なし（品目コード未登録）: {len(missing)} 件')
for x in missing: print('  行', x[0], x[1], x[2], f'払出 {x[3]}')
print(f'品目名の不一致（"?"以外の差）: {len(name_diff)} 件')
for x in name_diff: print('  行', x[0], x[1], f'案件管理表=[{x[2]}] SAP=[{x[3]}]')
print(f'残数ちょうど0（問題なし・参考）: {len(zero)} 件')
for x in zero: print('  行', x[0], x[1], x[2], f'払出 {x[3]}')

if out:
    from export import write_xlsx
    write_xlsx(out, req, sap, loc, sheet, start, end, req_path, sap_path, same_name)
    print('Excel出力:', out)
