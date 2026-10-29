"""払出後の在庫一覧Excelを作る（全品目 × 倉庫別、指定倉庫から払出数を差し引く）。

使い方:
  python3 after_stock.py <SAP在庫.xlsx> <払出リスト.tsv> <出力.xlsx> [払出元の保管場所コード=CX01]

払出リスト.tsv は「品目コード<TAB>数量」の行（ヘッダーなし、数量のカンマ可）。
在庫・払出後残数・合算はすべて数式で、元データ／払出リストのシートを参照する。
"""
import sys
from collections import OrderedDict
import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties

sap_path, req_path, out = sys.argv[1], sys.argv[2], sys.argv[3]
loc = sys.argv[4] if len(sys.argv) > 4 else 'CX01'

F = 'Calibri'
HFILL = PatternFill('solid', fgColor='305496'); TFILL = PatternFill('solid', fgColor='D9E1F2')
AFILL = PatternFill('solid', fgColor='E2EFDA')
THIN = Side(style='thin', color='BFBFBF'); BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = '#,##0;[Red]-#,##0'

src = openpyxl.load_workbook(sap_path)['XTab DS_1']
rows = [[c.value for c in r] for r in src.iter_rows(min_row=2, max_row=src.max_row) if r[1].value]
items = OrderedDict(); whs = OrderedDict()
for jan, code, txt, l, ln, q, a in rows:
    items.setdefault(code, (jan, txt)); whs.setdefault(l, ln)
if loc not in whs:
    sys.exit(f'払出元 {loc} が元データにありません')
others = [(l, n) for l, n in whs.items() if l != loc]

req = OrderedDict()
for line in open(req_path, encoding='utf-8'):
    p = line.rstrip('\n').split('\t')
    if len(p) < 2 or not p[0].strip(): continue
    code, q = p[0].strip(), int(str(p[-1]).replace(',', '').strip())
    req[code] = req.get(code, 0) + q
missing = [c for c in req if c not in items]

wb = openpyxl.Workbook()
ws = wb.active; ws.title = '払出後在庫一覧'
sd = wb.create_sheet('XTab DS_1'); rq = wb.create_sheet('払出リスト')


def header(sh, row, hdrs, widths):
    for i, (h, w) in enumerate(zip(hdrs, widths), 1):
        c = sh.cell(row, i, h); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HFILL; c.border = BD
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        sh.column_dimensions[c.column_letter].width = w
    sh.row_dimensions[row].height = 32


# 元データ（そのまま）
header(sd, 1, [src.cell(1, i).value for i in range(1, 8)], [18, 24, 52, 10, 12, 12, 14])
for r, row in enumerate(rows, 2):
    for c, v in enumerate(row, 1):
        cell = sd.cell(r, c, v); cell.font = Font(name=F)
        cell.number_format = NUM if c >= 6 else '@'
last = len(rows) + 1
sd.freeze_panes = 'A2'
S = lambda col: f"'XTab DS_1'!${col}$2:${col}${last}"

# 払出リスト（入力値は青字）
header(rq, 1, ['品目コード', '品目名（SAP）', '払出数'], [24, 60, 12])
for r, (code, q) in enumerate(req.items(), 2):
    rq.cell(r, 1, code).number_format = '@'
    rq.cell(r, 2, items[code][1] if code in items else '（元データに該当なし）')
    c = rq.cell(r, 3, q); c.number_format = NUM; c.font = Font(name=F, color='0000FF')
rlast = len(req) + 1
rq.cell(rlast + 1, 2, '合計').font = Font(name=F, bold=True)
rq.cell(rlast + 1, 3, f'=SUM(C2:C{rlast})').number_format = NUM
rq.freeze_panes = 'A2'
RQ = lambda col: f"払出リスト!${col}$2:${col}${rlast}"

# 一覧
locname = whs[loc]
ws['A1'] = f'払出後在庫一覧（払出元：{locname} {loc}）'; ws['A1'].font = Font(name=F, bold=True, size=14)
ws['A2'] = f'元データ：{sap_path.rsplit("/", 1)[-1]}　払出リスト：{len(req)}品目・合計 {sum(req.values()):,}'
ws['A3'] = f'払出は{locname}の在庫からのみ差し引く。払出後の残数が0は問題なし、マイナスのみ「不足」。'
for a in ('A2', 'A3'): ws[a].font = Font(name=F, color='595959')

H = 5
hdrs = ['JANコード(EAN/UPC)', '品目コード', '品目テキスト', f'{locname}\n払出前', '払出数', f'{locname}\n払出後']
hdrs += [n for _, n in others] + ['合算\n払出前', '合算\n払出後', '判定']
header(ws, H, hdrs, [16, 22, 50, 12, 10, 12] + [12] * len(others) + [12, 12, 8])
ncol = len(hdrs)
cB, cA = 7 + len(others), 8 + len(others)  # 合算 払出前 / 払出後
r = H + 1
for code, (jan, txt) in items.items():
    for c, v in enumerate([jan, code, txt], 1):
        ws.cell(r, c, v).number_format = '@'
    ws.cell(r, 4, f'=SUMIFS({S("F")},{S("B")},$B{r},{S("D")},"{loc}")')
    ws.cell(r, 5, f'=SUMIFS({RQ("C")},{RQ("A")},$B{r})')
    ws.cell(r, 6, f'=D{r}-E{r}')
    for i, (l, _) in enumerate(others):
        ws.cell(r, 7 + i, f'=SUMIFS({S("F")},{S("B")},$B{r},{S("D")},"{l}")')
    ws.cell(r, cB, f'=SUMIFS({S("F")},{S("B")},$B{r})')
    ws.cell(r, cA, f'=F{r}+' + '+'.join(ws.cell(r, 7 + i).coordinate for i in range(len(others))) if others else f'=F{r}')
    ws.cell(r, ncol, f'=IF(E{r}=0,"-",IF(F{r}<0,"不足","OK"))')
    for c in range(1, ncol + 1):
        cell = ws.cell(r, c); cell.border = BD; cell.font = Font(name=F)
        if c >= 4 and c < ncol: cell.number_format = NUM
    for c in (6, cA): ws.cell(r, c).fill = AFILL; ws.cell(r, c).font = Font(name=F, bold=True)
    ws.cell(r, ncol).alignment = Alignment(horizontal='center')
    r += 1
end = r - 1
ws.cell(r, 3, '合計').font = Font(name=F, bold=True)
for c in range(4, ncol):
    L = ws.cell(H, c).column_letter
    ws.cell(r, c, f'=SUM({L}{H+1}:{L}{end})').number_format = NUM
for c in range(1, ncol + 1):
    ws.cell(r, c).fill = TFILL; ws.cell(r, c).border = BD; ws.cell(r, c).font = Font(name=F, bold=True)
J = ws.cell(H, ncol).column_letter
s = r + 2
for i, (k, f) in enumerate([('払出対象の品目数', f'=COUNTIF({J}{H+1}:{J}{end},"<>-")'),
                            ('不足（マイナス）品目数', f'=COUNTIF({J}{H+1}:{J}{end},"不足")'),
                            (f'払出後に{locname}が0になる品目数', f'=COUNTIFS(E{H+1}:E{end},">0",F{H+1}:F{end},0)'),
                            ('払出リストの合計（照合用）', f'=払出リスト!C{rlast+1}')]):
    ws.cell(s + i, 2, k).font = Font(name=F, bold=True)
    c = ws.cell(s + i, 3, f); c.font = Font(name=F, bold=True); c.alignment = Alignment(horizontal='left'); c.number_format = NUM

area = f'A{H+1}:{J}{end}'
ws.conditional_formatting.add(area, FormulaRule(formula=[f'$F{H+1}<0'], fill=PatternFill('solid', fgColor='F8CBAD'), font=Font(color='9C0006', bold=True)))
ws.conditional_formatting.add(f'A{H+1}:E{end}', FormulaRule(formula=[f'$E{H+1}>0'], fill=PatternFill('solid', fgColor='FFF2CC')))
ws.freeze_panes = f'D{H+1}'; ws.auto_filter.ref = f'A{H}:{J}{end}'

wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(out)
print(out, f'品目 {len(items)} / 払出 {len(req)} 品目 合計 {sum(req.values())} / 元データにない品目 {missing}')
