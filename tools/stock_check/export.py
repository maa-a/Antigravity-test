"""check.py の結果を、品目ごとの在庫数・払出数・払出後残数の一覧Excelにする。

在庫数はSAP在庫シートを参照する SUMIFS、残数・判定・不足数も数式で計算する。
"""
import os
import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties

F = 'Calibri'
HFILL = PatternFill('solid', fgColor='305496')
TFILL = PatternFill('solid', fgColor='D9E1F2')
THIN = Side(style='thin', color='BFBFBF')
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = '#,##0;[Red]-#,##0'


def _header(ws, row, hdrs, widths):
    for i, (h, w) in enumerate(zip(hdrs, widths), 1):
        c = ws.cell(row, i, h)
        c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = HFILL; c.border = BD
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w


def write_xlsx(out, req, sap, loc, sheet, start, end, req_path, sap_path, same_name):
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = '払出チェック'
    locname = sap.loc[sap['保管場所'] == loc, '保管場所名'].iloc[0] if (sap['保管場所'] == loc).any() else loc
    ss = 'SAP在庫'

    # SAP在庫シート（指定倉庫の行のみ、元データのまま）
    st = wb.create_sheet(ss)
    cols = list(sap.columns)
    _header(st, 1, cols, [18, 24, 52, 10, 12, 12, 14])
    sub = sap[sap['保管場所'] == loc]
    for r, row in enumerate(sub.itertuples(index=False), 2):
        for c, v in enumerate(row, 1):
            cell = st.cell(r, c, v.item() if hasattr(v, 'item') else v)
            cell.font = Font(name=F); cell.border = BD
            if c >= 6: cell.number_format = NUM
            else: cell.number_format = '@'
    last = max(2, len(sub) + 1)
    st.freeze_panes = 'A2'; st.auto_filter.ref = f'A1:G{last}'
    rng = lambda col: f"{ss}!${col}$2:${col}${last}"

    # 見出し
    ws['A1'] = f'払出数量 在庫チェック（{locname} {loc}）'; ws['A1'].font = Font(name=F, bold=True, size=14)
    info = [('案件管理表', f'{os.path.basename(req_path)} ／ シート「{sheet}」 {start}～{end}行目（E列 品目コード・F列 品目名・H列 数量）'),
            ('在庫データ', f'{os.path.basename(sap_path)} の {locname}（{loc}）のみ'),
            ('判定ルール', '品目コードで照合。払出後の残数が0は問題なし（OK）、マイナスのみ「不足」。SAP品目名の「?」は表示できない文字のため同一商品とみなす。')]
    for i, (k, v) in enumerate(info, 2):
        ws.cell(i, 1, k).font = Font(name=F, bold=True); ws.cell(i, 2, v).font = Font(name=F)

    H = 7
    hdrs = ['案件管理表\n行', '品目コード', '品目名（案件管理表）', '品目名（SAP）', '品目名\n照合',
            f'在庫数\n（{locname}）', '払出数', '払出後\n残数', '判定', '不足数']
    _header(ws, H, hdrs, [10, 22, 50, 50, 9, 12, 10, 10, 8, 10])
    ws.row_dimensions[H].height = 32
    r = H + 1
    for code, e in req.items():
        m = sub[sub['品目コード'] == code]
        sapname = m['品目テキスト'].iloc[0] if len(m) else '（該当なし）'
        vals = [', '.join(map(str, e['rows'])), code, e['name'], sapname]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v).number_format = '@'
        ws.cell(r, 5, ('一致' if same_name(e['name'], sapname) else '要確認') if len(m) else '該当なし')
        ws.cell(r, 6, f'=SUMIFS({rng("F")},{rng("B")},$B{r})')
        ws.cell(r, 7, e['qty'])
        ws.cell(r, 8, f'=F{r}-G{r}')
        ws.cell(r, 9, f'=IF(H{r}<0,"不足","OK")')
        ws.cell(r, 10, f'=MAX(0,-H{r})')
        for c in range(1, 11):
            cell = ws.cell(r, c); cell.border = BD; cell.font = Font(name=F)
            if c >= 6 and c != 9: cell.number_format = NUM
            if c in (1, 5, 9): cell.alignment = Alignment(horizontal='center')
        ws.cell(r, 7).font = Font(name=F, color='0000FF')  # 案件管理表からの入力値
        r += 1
    end_r = r - 1

    # 合計・件数
    ws.cell(r, 3, '合計').font = Font(name=F, bold=True)
    for c, L in ((6, 'F'), (7, 'G'), (8, 'H'), (10, 'J')):
        ws.cell(r, c, f'=SUM({L}{H+1}:{L}{end_r})').number_format = NUM
    for c in range(1, 11):
        ws.cell(r, c).fill = TFILL; ws.cell(r, c).border = BD; ws.cell(r, c).font = Font(name=F, bold=True)
    s = r + 2
    summary = [('対象品目数', f'=COUNTA(B{H+1}:B{end_r})'),
               ('不足（マイナス）品目数', f'=COUNTIF(I{H+1}:I{end_r},"不足")'),
               ('残数0の品目数（問題なし）', f'=COUNTIF(H{H+1}:H{end_r},0)'),
               ('品目名 要確認・該当なし', f'=COUNTIF(E{H+1}:E{end_r},"<>一致")')]
    for i, (k, f) in enumerate(summary):
        ws.cell(s + i, 2, k).font = Font(name=F, bold=True)
        c = ws.cell(s + i, 3, f); c.font = Font(name=F, bold=True); c.alignment = Alignment(horizontal='left')
    ws.cell(s + len(summary) + 1, 2, '※払出数（青字）は案件管理表から転記。在庫数・残数・判定はSAP在庫シートを参照する数式。').font = Font(name=F, italic=True, color='595959')

    # 条件付き書式：不足は赤、残数0は灰色
    area = f'A{H+1}:J{end_r}'
    ws.conditional_formatting.add(area, FormulaRule(formula=[f'$H{H+1}<0'], fill=PatternFill('solid', fgColor='F8CBAD'), font=Font(color='9C0006', bold=True)))
    ws.conditional_formatting.add(area, FormulaRule(formula=[f'$H{H+1}=0'], fill=PatternFill('solid', fgColor='EDEDED')))
    ws.conditional_formatting.add(f'E{H+1}:E{end_r}', FormulaRule(formula=[f'$E{H+1}<>"一致"'], fill=PatternFill('solid', fgColor='FFF2CC')))
    ws.freeze_panes = f'C{H+1}'; ws.auto_filter.ref = f'A{H}:J{end_r}'

    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(out)
