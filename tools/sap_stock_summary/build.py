"""SAP在庫データ(XTab DS_1)から倉庫別シートと商品別_合算シートを追加する。

使い方: python3 build.py <元データ.xlsx> <出力.xlsx>
"""
import sys, openpyxl
from collections import OrderedDict
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.workbook.properties import CalcProperties

SRC, OUT = sys.argv[1], sys.argv[2]
wb=openpyxl.load_workbook(SRC)
src=wb['XTab DS_1']; S="'XTab DS_1'"
QF='#,##0 "PC";-#,##0 "PC"'; AF='#,##0 "JPY";-#,##0 "JPY"'
last=src.max_row
rows=[[c.value for c in r] for r in src.iter_rows(min_row=2,max_row=last)]

items=OrderedDict(); whs=OrderedDict()
for jan,code,txt,loc,locname,q,a in rows:
    items.setdefault(code,{'jan':jan,'txt':txt,'locs':[]})['locs'].append(loc)
    whs.setdefault(loc,locname)

hfont=Font(name='Calibri',bold=True,color='FFFFFF'); hfill=PatternFill('solid',fgColor='305496')
tfill=PatternFill('solid',fgColor='D9E1F2'); sfill=PatternFill('solid',fgColor='EDEDED')
dupfill=PatternFill('solid',fgColor='FFF2CC')
thin=Side(style='thin',color='BFBFBF'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
bold=Font(name='Calibri',bold=True)
R=lambda col:f"{S}!${col}$2:${col}${last}"
KBN=lambda r:f'=IF(COUNTIFS({R("B")},$B{r})>=2,"重複","単独")'

def header(ws,hdrs,widths):
    for i,(h,w) in enumerate(zip(hdrs,widths),1):
        c=ws.cell(1,i,h); c.font=hfont; c.fill=hfill; c.border=bd
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.column_dimensions[c.column_letter].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref=f"A1:{ws.cell(1,len(hdrs)).column_letter}{ws.max_row}"

def totals(ws,r,kcol,sum_cols,fmts,ncols):
    """合計行 + 区分別小計（重複/単独）"""
    end=r-1; K=ws.cell(1,kcol).column_letter
    for label,cond,fill in [('合計',None,tfill),('うち 重複',"重複",sfill),('うち 単独',"単独",sfill)]:
        ws.cell(r,3,label).font=bold
        for col,f in zip(sum_cols,fmts):
            L=ws.cell(1,col).column_letter
            fx=f"=SUM({L}2:{L}{end})" if cond is None else f'=SUMIFS({L}2:{L}{end},{K}2:{K}{end},"{cond}")'
            c=ws.cell(r,col,fx); c.number_format=f; c.font=bold
        for i in range(1,ncols+1): ws.cell(r,i).fill=fill; ws.cell(r,i).border=bd
        r+=1
    return r

def dup_highlight(ws,r,ncols,code):
    if len(set(items[code]['locs']))>=2:
        for i in range(1,ncols+1): ws.cell(r,i).fill=dupfill

# 倉庫別シート（重複・単独すべて）
for loc,locname in whs.items():
    ws=wb.create_sheet(f"倉庫別_{locname}")
    hdrs=['JANコード(EAN/UPC)','品目コード','品目テキスト','保管場所','保管場所名','区分','在庫数量','在庫金額']
    header(ws,hdrs,[18,24,52,10,12,8,14,16])
    r=2
    for code,v in items.items():
        if loc not in v['locs']: continue
        for i,x in enumerate([v['jan'],code,v['txt'],loc,locname],1): ws.cell(r,i,x).number_format='@'
        ws.cell(r,6,KBN(r)).alignment=Alignment(horizontal='center')
        ws.cell(r,7,f'=SUMIFS({R("F")},{R("B")},$B{r},{R("D")},$D{r})').number_format=QF
        ws.cell(r,8,f'=SUMIFS({R("G")},{R("B")},$B{r},{R("D")},$D{r})').number_format=AF
        for i in range(1,9): ws.cell(r,i).border=bd
        dup_highlight(ws,r,8,code); r+=1
    ws.auto_filter.ref=f"A1:H{r-1}"
    totals(ws,r,6,[7,8],[QF,AF],8)

# 合算シート（全商品）
ws=wb.create_sheet('商品別_合算')
wl=list(whs.items())
hdrs=['JANコード(EAN/UPC)','品目コード','品目テキスト','登録倉庫数','区分']
for loc,n in wl: hdrs+=[f'{n} 数量',f'{n} 金額']
hdrs+=['合算 在庫数量','合算 在庫金額(原価)']
header(ws,hdrs,[18,24,52,10,8]+[13,14]*len(wl)+[15,18])
ncol=len(hdrs); qcol=ncol-1; acol=ncol
for r,(code,v) in enumerate(items.items(),2):
    for i,x in enumerate([v['jan'],code,v['txt']],1): ws.cell(r,i,x).number_format='@'
    ws.cell(r,4,f'=COUNTIFS({R("B")},$B{r})').alignment=Alignment(horizontal='center')
    ws.cell(r,5,f'=IF(D{r}>=2,"重複","単独")').alignment=Alignment(horizontal='center')
    col=6
    for loc,n in wl:
        ws.cell(r,col,f'=SUMIFS({R("F")},{R("B")},$B{r},{R("D")},"{loc}")').number_format=QF
        ws.cell(r,col+1,f'=SUMIFS({R("G")},{R("B")},$B{r},{R("D")},"{loc}")').number_format=AF
        col+=2
    ws.cell(r,qcol,f'=SUMIFS({R("F")},{R("B")},$B{r})').number_format=QF
    ws.cell(r,acol,f'=SUMIFS({R("G")},{R("B")},$B{r})').number_format=AF
    for i in range(1,ncol+1): ws.cell(r,i).border=bd
    for c in (qcol,acol): ws.cell(r,c).font=bold
    dup_highlight(ws,r,ncol,code)
tr=len(items)+2
ws.auto_filter.ref=f"A1:{ws.cell(1,ncol).column_letter}{tr-1}"
nr=totals(ws,tr,5,list(range(6,ncol+1)),[QF,AF]*(len(wl)+1),ncol)
ws.cell(nr+1,1,'※元データ「XTab DS_1」の全品目を掲載。2つ以上の倉庫(保管場所)に登録されている品目を「重複」（黄色）、1倉庫のみを「単独」と表示。数量・金額はSUMIFSで元データから集計（元データ変更で自動再計算）。').font=Font(name='Calibri',italic=True,color='595959')
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print(OUT, len(items), sum(len(set(v['locs']))>=2 for v in items.values()), wl)
